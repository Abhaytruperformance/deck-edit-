"""XLSX renderer: same Artifact Contract blocks, one handler per block type,
written sequentially down a single worksheet since a spreadsheet has no
concept of "slides".
"""
import io

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from app.db import get_supabase

_ASSET_BUCKET = "deliverable-assets"
_TITLE_FONT = Font(size=16, bold=True)
_HEADING_FONT = Font(size=13, bold=True)
_BOLD = Font(bold=True)


def _handle_title_slide(ws: Worksheet, content: dict, row: int) -> int:
    ws.cell(row=row, column=1, value=content["headline"]).font = _TITLE_FONT
    ws.cell(row=row + 1, column=1, value=content["subhead"])
    return row + 3


def _handle_text_block(ws: Worksheet, content: dict, row: int) -> int:
    ws.cell(row=row, column=1, value=content["heading"]).font = _HEADING_FONT
    cell = ws.cell(row=row + 1, column=1, value=content["body"])
    cell.alignment = Alignment(wrap_text=True)
    return row + 3


def _handle_bullet_list(ws: Worksheet, content: dict, row: int) -> int:
    ws.cell(row=row, column=1, value=content["heading"]).font = _HEADING_FONT
    row += 1
    for item in content["items"]:
        ws.cell(row=row, column=1, value=f"- {item}")
        row += 1
    return row + 1


def _handle_kpi_grid(ws: Worksheet, content: dict, row: int) -> int:
    ws.cell(row=row, column=1, value=content["title"]).font = _HEADING_FONT
    row += 1
    for i, item in enumerate(content["items"], start=1):
        ws.cell(row=row, column=1, value=item["label"]).font = _BOLD
        ws.cell(row=row, column=2, value=item["value"])
        ws.cell(row=row, column=3, value=item["unit"])
        if item.get("change_pct") is not None:
            ws.cell(row=row, column=4, value=f"{item['change_pct']:+.1f}%")
        row += 1
    return row + 1


def _handle_comparison_table(ws: Worksheet, content: dict, row: int) -> int:
    for c, header in enumerate(content["headers"], start=1):
        ws.cell(row=row, column=c, value=header).font = _BOLD
    row += 1
    for data_row in content["rows"]:
        for c, cell_value in enumerate(data_row, start=1):
            ws.cell(row=row, column=c, value=cell_value)
        row += 1
    return row + 1


_CHART_CLASSES = {"bar": BarChart, "line": LineChart, "pie": PieChart}


def _handle_chart(ws: Worksheet, content: dict, row: int) -> int:
    ws.cell(row=row, column=1, value=content["title"]).font = _HEADING_FONT
    data_start_row = row + 1

    ws.cell(row=data_start_row, column=1, value="Category")
    for s, series in enumerate(content["series"], start=2):
        ws.cell(row=data_start_row, column=s, value=series["name"])
    for i, label in enumerate(content["labels"], start=1):
        ws.cell(row=data_start_row + i, column=1, value=label)
    for s, series in enumerate(content["series"], start=2):
        for i, value in enumerate(series["values"], start=1):
            ws.cell(row=data_start_row + i, column=s, value=value)

    n_rows = len(content["labels"])
    n_series = len(content["series"])
    chart = _CHART_CLASSES[content["chart_type"]]()
    chart.title = content["title"]
    data = Reference(ws, min_col=2, max_col=1 + n_series, min_row=data_start_row, max_row=data_start_row + n_rows)
    cats = Reference(ws, min_col=1, min_row=data_start_row + 1, max_row=data_start_row + n_rows)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    anchor_row = data_start_row + n_rows + 2
    ws.add_chart(chart, f"A{anchor_row}")
    return anchor_row + 16  # chart occupies visual space even though rows below are unused


def _handle_image_block(ws: Worksheet, content: dict, row: int) -> int:
    try:
        image_bytes = get_supabase().storage.from_(_ASSET_BUCKET).download(content["image_ref"])
        ws.add_image(XLImage(io.BytesIO(image_bytes)), f"A{row}")
        row += 16
    except Exception:
        # ponytail: same tradeoff as the PPTX renderer - skip a broken image
        # ref rather than failing the whole export.
        pass
    ws.cell(row=row, column=1, value=content["caption"])
    return row + 2


_HANDLERS = {
    "title_slide": _handle_title_slide,
    "text_block": _handle_text_block,
    "bullet_list": _handle_bullet_list,
    "kpi_grid": _handle_kpi_grid,
    "comparison_table": _handle_comparison_table,
    "chart": _handle_chart,
    "image_block": _handle_image_block,
}


def render(blocks: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    row = 1
    for block in blocks:
        if not block.get("enabled", True):
            continue
        row = _HANDLERS[block["type"]](ws, block["content"], row)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
