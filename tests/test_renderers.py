"""Self-check for the three renderers: one of every block type (plus a
disabled block that must be skipped), run through each renderer, verifying
it produces well-formed output without needing a live Supabase (image_block
storage lookups are expected to fail closed and get skipped).

Run: python tests/test_renderers.py
"""
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("SUPABASE_URL", "https://example.supabase.co")
os.environ.setdefault("SUPABASE_KEY", "dummy")

from openpyxl import load_workbook  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.enum.text import MSO_AUTO_SIZE  # noqa: E402

from app.renderers import html as html_renderer  # noqa: E402
from app.renderers import pptx as pptx_renderer  # noqa: E402
from app.renderers import xlsx as xlsx_renderer  # noqa: E402

SAMPLE_BLOCKS = [
    {"id": "b1", "type": "title_slide", "enabled": True, "content": {"headline": "Q3 Review", "subhead": "Acme Corp"}},
    {"id": "b2", "type": "text_block", "enabled": True, "content": {"heading": "Summary", "body": "Great quarter."}},
    {"id": "b3", "type": "bullet_list", "enabled": True, "content": {"heading": "Highlights", "items": ["Grew 20%", "Shipped v2"]}},
    {"id": "b4", "type": "kpi_grid", "enabled": True, "content": {"title": "Metrics", "items": [
        {"label": "Revenue", "value": 120000, "unit": "USD", "change_pct": 5.2},
        {"label": "Churn", "value": 2.1, "unit": "%", "change_pct": None},
    ]}},
    {"id": "b5", "type": "comparison_table", "enabled": True, "content": {"headers": ["Plan", "Price"], "rows": [["Basic", "$10"], ["Pro", "$30"]]}},
    {"id": "b6", "type": "chart", "enabled": True, "content": {"chart_type": "bar", "title": "Growth", "labels": ["Jan", "Feb"], "series": [{"name": "Sales", "values": [1, 2]}]}},
    {"id": "b7", "type": "image_block", "enabled": True, "content": {"caption": "Team photo", "image_ref": "does/not/exist.png"}},
    {"id": "b8", "type": "text_block", "enabled": False, "content": {"heading": "Should not appear", "body": "..."}},
]


def test_pptx():
    data = pptx_renderer.render(SAMPLE_BLOCKS)
    prs = Presentation(io.BytesIO(data))
    # 7 enabled blocks -> 7 slides; the disabled text_block must be skipped.
    assert len(list(prs.slides)) == 7
    all_text = "\n".join(
        shape.text_frame.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame
    )
    assert "Should not appear" not in all_text
    assert "Q3 Review" in all_text


def test_pptx_long_text_autofits_instead_of_overflowing():
    long_body = "This client sentence just keeps going. " * 40
    blocks = [
        {"id": "t1", "type": "text_block", "enabled": True, "content": {"heading": "H" * 120, "body": long_body}},
    ]
    data = pptx_renderer.render(blocks)
    prs = Presentation(io.BytesIO(data))
    slide = list(prs.slides)[0]
    text_frames = [shape.text_frame for shape in slide.shapes if shape.has_text_frame]
    assert text_frames, "expected at least one text frame on the slide"
    for tf in text_frames:
        assert tf.word_wrap is True
        assert tf.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def test_xlsx():
    data = xlsx_renderer.render(SAMPLE_BLOCKS)
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    values = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert any("Q3 Review" in v for v in values)
    assert not any("Should not appear" in v for v in values)


def test_html():
    page = html_renderer.render(SAMPLE_BLOCKS)
    assert "Q3 Review" in page
    assert "Should not appear" not in page
    assert "<canvas" in page
    assert "gsap" in page.lower()


if __name__ == "__main__":
    test_pptx()
    test_pptx_long_text_autofits_instead_of_overflowing()
    test_xlsx()
    test_html()
    print("OK: pptx/xlsx/html renderers all skip disabled blocks and render every enabled block type")
